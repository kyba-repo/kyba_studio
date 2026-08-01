from __future__ import annotations

import os
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"

from pathlib import Path  # noqa: E402
from typing import Any, Optional  # noqa: E402

from fastapi import FastAPI, HTTPException, Request, BackgroundTasks  # noqa: E402
from pydantic import BaseModel  # noqa: E402
import requests  # noqa: E402

from langchain_chroma import Chroma  # noqa: E402
from langchain_community.document_loaders import PyPDFLoader, TextLoader, Docx2txtLoader, CSVLoader  # noqa: E402
from langchain_huggingface import HuggingFaceEmbeddings  # noqa: E402
from langchain_text_splitters import RecursiveCharacterTextSplitter  # noqa: E402
from langchain_core.documents import Document  # noqa: E402

USER_KYBA_DIR = Path.home() / ".kyba"
KNOWLEDGE_BASE_DIR = USER_KYBA_DIR / "knowledge"
DB_BASE_DIR = USER_KYBA_DIR / "chroma_db"
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".json", ".csv", ".html", ".htm", ".pdf", ".docx", ".xlsx"}

app = FastAPI(title="Kyba RAG Server", version="0.1.0")


class ChatRequest(BaseModel):
    question: str
    options: Optional[dict] = None
    system_prompt: Optional[str] = None
    history: Optional[list[dict[str, Any]]] = None
    profile_id: str = "default"
    model: str = "gemma4:e2b"
    images: Optional[list[str]] = None
    documents: Optional[list[dict[str, str]]] = None
    reasoning_effort: Optional[str] = "medium"


class IngestRequest(BaseModel):
    path: Optional[str] = None
    profile_id: str = "default"


vectorstores: dict[str, Chroma] = {}
retrievers: dict[str, Any] = {}
global_embeddings = None
whisper_model = None


def _ensure_whisper() -> Any:
    global whisper_model
    if whisper_model is None:
        try:
            from faster_whisper import WhisperModel
            import os
            
            # Search for locally packaged model
            cwd = os.getcwd()
            local_model_path = os.path.join(cwd, "models", "faster-whisper-tiny")
            
            if os.path.exists(local_model_path):
                print(f"[Kyba] Using local Whisper model at {local_model_path}")
                whisper_model = WhisperModel(local_model_path, device="cpu", compute_type="int8", local_files_only=True)
            else:
                print("[Kyba] Local model not found, using HuggingFace cache...")
                whisper_model = WhisperModel("tiny", device="cpu", compute_type="int8")
        except ImportError:
            raise RuntimeError("faster-whisper is not installed.")
    return whisper_model


llm_model: str = os.getenv("OLLAMA_MODEL", "gemma4:e2b")
llm_base_url: str = os.getenv("OLLAMA_HOST", "http://127.0.0.1:11434")


def _ensure_runtime(profile_id: str = "default") -> tuple[Chroma, Any]:
    global global_embeddings
    if global_embeddings is None:
        try:
            global_embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
        except Exception:
            # Fallback to offline mode if there is no internet (assumes it was downloaded before)
            os.environ["HF_HUB_OFFLINE"] = "1"
            os.environ["TRANSFORMERS_OFFLINE"] = "1"
            global_embeddings = HuggingFaceEmbeddings(
                model_name="all-MiniLM-L6-v2",
                model_kwargs={"local_files_only": True}
            )
    
    if profile_id not in vectorstores:
        db_path = DB_BASE_DIR / profile_id
        db_path.mkdir(parents=True, exist_ok=True)
        vectorstores[profile_id] = Chroma(
            collection_name=f"kyba_knowledge_{profile_id}",
            embedding_function=global_embeddings,
            persist_directory=str(db_path),
        )
        retrievers[profile_id] = vectorstores[profile_id].as_retriever(search_kwargs={"k": 4})
    
    return vectorstores[profile_id], retrievers[profile_id]


def _load_single_document(file_path: Path) -> list[Any]:
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            loader = PyPDFLoader(str(file_path))
            return loader.load()
        elif suffix == ".docx":
            loader = Docx2txtLoader(str(file_path))
            return loader.load()
        elif suffix == ".csv":
            loader = CSVLoader(str(file_path), encoding="utf-8")
            return loader.load()
        elif suffix == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            text_content = ""
            for sheet in wb.worksheets:
                text_content += f"\n--- Sheet: {sheet.title} ---\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                    text_content += row_text + "\n"
            return [Document(page_content=text_content, metadata={"source": str(file_path)})]

        loader = TextLoader(str(file_path), encoding="utf-8")
        return loader.load()
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"Could not load {file_path}: {exc}") from exc


def _load_documents(path: Optional[str] = None, profile_id: str = "default") -> list[Any]:
    default_target = KNOWLEDGE_BASE_DIR / profile_id
    target_path = Path(path).resolve() if path else default_target
    if not target_path.exists():
        return []

    if target_path.is_file():
        return _load_single_document(target_path)

    documents: list[Any] = []
    for file_path in sorted(target_path.rglob("*")):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            documents.extend(_load_single_document(file_path))
        except Exception as e:
            print(f"Skipping {file_path} due to error: {e}")
    return documents


def _ingest_documents(path: Optional[str] = None, profile_id: str = "default") -> int:
    vs, _ = _ensure_runtime(profile_id)
    if vs is None:
        raise RuntimeError(f"Could not initialize vector database for profile {profile_id}.")

    docs = _load_documents(path, profile_id)
    if not docs:
        return 0

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)
    chunks = splitter.split_documents(docs)
    if not chunks:
        return 0

    # Batch add documents to reduce embedding calls and I/O.
    batch_size = int(os.getenv("INGEST_BATCH_SIZE", "128"))
    persist_every = int(os.getenv("INGEST_PERSIST_EVERY", "10"))
    indexed = 0
    failed = 0
    batch_count = 0
    last_error = None
    for i in range(0, len(chunks), batch_size):
        batch = chunks[i:i + batch_size]
        try:
            vs.add_documents(batch)
            indexed += len(batch)
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            # Fallback: try adding documents one-by-one to isolate failures
            print(f"Batch add failed (size={len(batch)}): {exc}. Falling back to per-chunk add.")
            for chunk in batch:
                try:
                    vs.add_documents([chunk])
                    indexed += 1
                except Exception as e:  # noqa: BLE001
                    failed += 1
                    try:
                        src = chunk.metadata.get("source", "no source")
                    except Exception:
                        src = "no source"
                    print(f"Failed chunk: {src} -> {e}")

        batch_count += 1
        # Persist periodically to avoid large in-memory state and I/O spikes
        if batch_count % persist_every == 0:
            try:
                vs.persist()
            except Exception as e:  # noqa: BLE001
                print(f"Persist failed in batch {batch_count}: {e}")

    if indexed == 0:
        raise RuntimeError(
            f"Could not index any chunks. Last error: {last_error if last_error else 'no details'}"
        )

    # Final persist to ensure everything is written to disk
    try:
        vs.persist()
    except Exception as e:  # noqa: BLE001
        print(f"Final persist failed: {e}")

    if failed:
        print(f"Indexed {indexed} chunks and skipped {failed} problematic chunks.")
    return indexed


def _format_context(docs: list[Any]) -> str:
    if not docs:
        return "No relevant context available."
    parts = []
    for idx, doc in enumerate(docs, start=1):
        source = doc.metadata.get("source", "no source")
        parts.append(f"[{idx}] Source: {source}\n{doc.page_content}")
    return "\n\n".join(parts)


@app.on_event("startup")
def startup() -> None:
    # Ensure default profile runs on startup
    _ensure_runtime("default")
    try:
        if not os.path.exists(DB_BASE_DIR):
            os.makedirs(DB_BASE_DIR, exist_ok=True)
        # Background ingest for default profile
        import threading
        threading.Thread(target=_ingest_documents, args=(None, "default"), daemon=True).start()
    except Exception:
        pass


@app.get("/")
def health() -> dict[str, Any]:
    return {"status": "ok", "model": "gemma4:e2b", "knowledge_dir": str(KNOWLEDGE_BASE_DIR)}


@app.post("/transcribe")
async def transcribe(request: Request) -> dict[str, Any]:
    try:
        model = _ensure_whisper()
        body = await request.body()
        if not body:
            raise HTTPException(status_code=400, detail="No audio provided")
        
        import tempfile
        with tempfile.NamedTemporaryFile(delete=False, suffix=".webm") as tmp:
            tmp.write(body)
            tmp_path = tmp.name
        
        try:
            segments, info = model.transcribe(tmp_path, language="es")
            text = "".join([segment.text for segment in segments])
            return {"text": text.strip()}
        finally:
            import os
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/ingest")
def ingest(payload: IngestRequest, background_tasks: BackgroundTasks) -> dict[str, Any]:
    try:
        background_tasks.add_task(_ingest_documents, payload.path, payload.profile_id)
        return {"status": "ok", "message": "Background indexing started"}
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@app.post("/chat")
def chat(payload: ChatRequest) -> dict[str, Any]:
    if not payload.question or not payload.question.strip():
        raise HTTPException(status_code=400, detail="You must send a question.")

    try:
        _, retriever = _ensure_runtime(payload.profile_id)
        docs = retriever.invoke(payload.question)  # type: ignore[union-attr]
        context = _format_context(docs)
        # Allow overriding the system prompt from the request; otherwise use the default template
        if payload.system_prompt is not None:
            system_prompt = payload.system_prompt.strip()
        else:
            system_prompt = (
                "You are an expert Python code generation assistant. Always answer in Spanish. Strictly follow these rules:\n\n"
                "Role and objective: Act as a senior Python developer who writes clean, tested, and maintainable code. Prioritize clarity, simplicity, and reasonable performance.\n"
                "Execution Environment: You are running locally on the user's machine with FULL administrative privileges, FULL internet access, and NO network restrictions. If the user asks you to run a command, install a package (e.g., pip install), or fetch data from the internet, DO IT using your tools. Do not refuse claiming network restrictions or lack of permissions.\n"
                "Output: Provide ONLY the requested result (code and/or tests). Do not show your \"thinking\", internal processes, or intermediate steps.\n"
                "Code style: Follow PEP 8 and use static typing (type hints) when possible. Write small functions with a single responsibility.\n"
                "Include clear docstrings in Google or NumPy format for public functions. Avoid unnecessary dependencies; use the standard library when sufficient.\n"
                "Security and robustness: Validate inputs and handle errors with specific exceptions. Avoid arbitrary code execution unless explicitly requested and confirmed by the user.\n"
                "Tests and verification: When generating modules, add at least one unit test with pytest that covers normal and edge cases. Include usage examples (if __name__ == \"__main__\": block if applicable).\n"
                "Delivery format: If the user requests a module: return only the complete .py file. If multiple files are required, deliver one file per section and precede each with a comment indicating the file name.\n"
                "Conversation constraints: Do not dig unnecessarily; if information is missing, ask only what is essential in a short and concrete question. Do not include long explanations unless the user asks for a review or comments.\n"
            )
        
        # Limit the length of the "thinking" if it's a reasoning model
        effort = payload.reasoning_effort or "medium"
        if effort == "low":
            system_prompt += "\nRULE: You can briefly analyze the question before answering, but get straight to the point quickly and without rambling."
        elif effort == "medium":
            system_prompt += "\nRULE: Take a moment to organize your ideas and think logically, but keep your analysis concise before giving the final answer."
        elif effort == "high":
            system_prompt += "\nRULE: You are free to think in detail, explore options and reason step by step as much as you need before delivering the best possible answer."
            
        import base64
        import tempfile
        
        extra_context = ""
        if getattr(payload, "documents", None):
            for doc in payload.documents:
                try:
                    filename = doc.get("filename", "documento.pdf")
                    content_b64 = doc.get("content", "")
                    if content_b64.startswith("data:"):
                        content_b64 = content_b64.split(",")[1]
                    
                    file_data = base64.b64decode(content_b64)
                    suffix = os.path.splitext(filename)[1].lower()
                    
                    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp_file:
                        tmp_file.write(file_data)
                        tmp_path = tmp_file.name
                        
                    docs = _load_single_document(Path(tmp_path))
                    extra_context += f"\n--- Content of {filename} ---\n"
                    for d in docs:
                        extra_context += d.page_content + "\n"
                        
                    os.unlink(tmp_path)
                except Exception as e:
                    print(f"Error processing document: {e}")
                    
        user_prompt = f"Context:\n{context}\n{extra_context}\nQuestion: {payload.question}"
        options = payload.options if getattr(payload, 'options', None) else {"temperature": 0.2}
        options["num_ctx"] = 8192
        
        messages = [{"role": "system", "content": system_prompt}]
        if payload.history:
            # Keep only the last 6 messages to avoid overflowing the model's context
            messages.extend(payload.history[-6:])
            
        user_msg = {"role": "user", "content": user_prompt}
        if getattr(payload, "images", None):
            clean_images = []
            for img in payload.images:
                if img.startswith("data:"):
                    clean_images.append(img.split(",")[1])
                else:
                    clean_images.append(img)
            user_msg["images"] = clean_images
            
        messages.append(user_msg)
        
        # Unload any other model from VRAM that is not the requested one using /api/ps
        try:
            ps_response = requests.get(f"{llm_base_url}/api/ps", timeout=5)
            if ps_response.ok:
                loaded_models = ps_response.json().get("models", [])
                for m in loaded_models:
                    m_name = m.get("name", "")
                    if m_name and m_name != payload.model:
                        print(f"[Kyba][backend] Unloading previous model from VRAM: {m_name}")
                        requests.post(
                            f"{llm_base_url}/api/generate",
                            json={"model": m_name, "keep_alive": 0},
                            timeout=10
                        )
        except Exception as e:
            print(f"Error checking/unloading previous models: {e}")
        
        response = requests.post(
            f"{llm_base_url}/api/chat",
            json={
                "model": payload.model,
                "messages": messages,
                "stream": False,
                "options": options,
                "keep_alive": "15m",
            },
            timeout=600,
        )
        response.raise_for_status()
        data = response.json()
        
        message = data.get("message", {})
        answer = message.get("content", "") or data.get("response", "")
        
        # Filter <think> tags from reasoning models (like DeepSeek-R1)
        import re
        answer = re.sub(r"<think>.*?</think>", "", answer, flags=re.DOTALL).strip()
        
        return {
            "answer": answer,
            "sources": [doc.metadata.get("source", "no source") for doc in docs],
        }
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=str(exc)) from exc
